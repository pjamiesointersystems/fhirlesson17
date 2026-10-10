"""Generate the Lesson 17 Instructor Guide as a Word document."""

from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
import os

doc = Document()

# ── Styles ──────────────────────────────────────────────────────────────
style = doc.styles["Normal"]
font = style.font
font.name = "Calibri"
font.size = Pt(11)

# ── Title Page ──────────────────────────────────────────────────────────
doc.add_paragraph()
doc.add_paragraph()
title = doc.add_paragraph()
title.alignment = WD_ALIGN_PARAGRAPH.CENTER
run = title.add_run("Lesson 17: Using LLMs with FHIR")
run.bold = True
run.font.size = Pt(28)
run.font.color.rgb = RGBColor(0, 51, 102)

subtitle = doc.add_paragraph()
subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
run = subtitle.add_run("Instructor Guide")
run.bold = True
run.font.size = Pt(20)
run.font.color.rgb = RGBColor(0, 102, 153)

doc.add_paragraph()

info = doc.add_paragraph()
info.alignment = WD_ALIGN_PARAGRAPH.CENTER
run = info.add_run("Vanderbilt FHIR Training Course\n")
run.font.size = Pt(14)
run = info.add_run("Author: Patrick W. Jamieson, M.D.\n")
run.font.size = Pt(12)
run = info.add_run("Four-Part AI+FHIR Series — Lesson 1 of 4")
run.font.size = Pt(12)
run.italic = True

doc.add_page_break()

# ── Table of Contents placeholder ───────────────────────────────────────
doc.add_heading("Table of Contents", level=1)
toc_items = [
    "1. Lesson Overview and Goals",
    "2. Series Context — Where This Lesson Fits",
    "3. Pre-Class Preparation Checklist",
    "4. Phase 1: AI Foundations Lecture Guide",
    "5. Phase 2: Exercise Walkthroughs",
    "    5.1 Exercise 1 — LLM Hello World",
    "    5.2 Exercise 2 — FHIR Bundle Exploration",
    "    5.3 Exercise 3 — Patient Summarization Pipeline",
    "    5.4 Exercise 4 — Prompt Engineering Comparison",
    "6. Phase 3: Safety, Governance, and Series Bridge",
    "7. Timing and Pacing Guide",
    "8. Common Issues and Troubleshooting",
    "9. Discussion Questions and Assessment Ideas",
    "10. Appendix: File Reference",
]
for item in toc_items:
    p = doc.add_paragraph(item)
    p.paragraph_format.space_after = Pt(2)

doc.add_page_break()

# ── 1. Lesson Overview ──────────────────────────────────────────────────
doc.add_heading("1. Lesson Overview and Goals", level=1)

doc.add_paragraph(
    "Lesson 17 introduces Large Language Models (LLMs) to students who have completed "
    "the prior FHIR lessons but have no AI/ML background. The lesson uses a local LLM "
    "(Ollama) to keep patient data private and connects every AI concept to a healthcare "
    "analogy the students already understand."
)

doc.add_paragraph(
    "By the end of this lesson, students will have:"
)

objectives = [
    "Understood 9 foundational AI concepts through healthcare analogies",
    "Verified their local LLM setup with a simple healthcare prompt (Exercise 1)",
    "Explored a real FHIR patient bundle and seen why it doesn't fit in a context window (Exercise 2)",
    "Built a working patient summarization pipeline that reduces ~100K tokens to ~1K (Exercise 3)",
    "Compared prompt engineering strategies side-by-side and observed temperature effects (Exercise 4)",
    "Discussed AI safety principles including the 'golden rule' of validation",
]
for obj in objectives:
    doc.add_paragraph(obj, style="List Bullet")

doc.add_heading("What This Lesson Does NOT Cover", level=2)
doc.add_paragraph(
    "Agentic AI, Model Context Protocol (MCP), and Retrieval-Augmented Generation (RAG) "
    "are explicitly out of scope. These topics are covered in Parts One, Two, and the RAG Lab "
    "respectively. If students ask about agents or RAG, acknowledge the topics and note that "
    "they are taught in the subsequent lessons — this builds anticipation without derailing the "
    "current lesson."
)

# ── 2. Series Context ───────────────────────────────────────────────────
doc.add_heading("2. Series Context — Where This Lesson Fits", level=1)

doc.add_paragraph(
    "Lesson 17 is the first lesson in a four-part AI+FHIR series. It is essential that "
    "students understand this context — they are not learning a standalone topic but building "
    "the foundation for increasingly sophisticated applications."
)

table = doc.add_table(rows=5, cols=3)
table.style = "Light Grid Accent 1"
table.alignment = WD_TABLE_ALIGNMENT.CENTER
headers = ["#", "Lesson", "Focus"]
for i, h in enumerate(headers):
    table.rows[0].cells[i].text = h
    for p in table.rows[0].cells[i].paragraphs:
        for r in p.runs:
            r.bold = True

series_data = [
    ("1", "Lesson 17 (this lesson)", "AI foundations, hands-on exercises, safety framing"),
    ("2", "Agentic AI — Part One", "Agent architecture, MCP, tool-use, structured validation, audit logging"),
    ("3", "Agentic AI — Part Two", "FHIR as knowledge layer, subagents, human review checkpoints"),
    ("4", "RAG Lab", "Vector stores, embeddings for retrieval, hybrid search, Graph RAG"),
]
for row_idx, (num, lesson, focus) in enumerate(series_data, start=1):
    table.rows[row_idx].cells[0].text = num
    table.rows[row_idx].cells[1].text = lesson
    table.rows[row_idx].cells[2].text = focus

doc.add_paragraph()
doc.add_paragraph(
    "Key message to students: Every concept introduced today reappears in later lessons. "
    "The summarization pipeline becomes a subagent. The feature extraction feeds a RAG pipeline. "
    "The prompt engineering techniques apply to agent prompts. Safety is a design constraint "
    "throughout."
)

# ── 3. Pre-Class Preparation ────────────────────────────────────────────
doc.add_heading("3. Pre-Class Preparation Checklist", level=1)

doc.add_paragraph(
    "Complete these steps before class to ensure a smooth lesson. Most setup issues occur "
    "with Ollama installation or FHIR server connectivity."
)

doc.add_heading("Instructor Setup", level=2)
instructor_prep = [
    "Install Ollama and pull the llama3.2:3b model (ollama pull llama3.2:3b)",
    "Verify Ollama is running: ollama serve, then test with ollama run llama3.2:3b",
    "Verify FHIR server is running with Synthea data: curl http://127.0.0.1:8080/csp/healthshare/demo/fhir/r4/Patient/2",
    "Run all four exercises once to confirm they work and to familiarize yourself with the output",
    "Review slides/ai_foundations.md and slides/safety_governance.md for the lecture portions",
    "If using Angular for demos, ensure the Angular front end is configured (separate from this lesson)",
    "Prepare .env file: copy .env.example to .env and fill in FHIR_PASSWORD",
]
for item in instructor_prep:
    doc.add_paragraph(item, style="List Bullet")

doc.add_heading("Student Pre-Requisites to Communicate", level=2)
student_prereqs = [
    "Python 3.11+ installed with venv support",
    "Prior FHIR lessons completed (Patient, Observation, Condition, Bundle, FHIRPath)",
    "InterSystems IRIS for Health running locally with Synthea data loaded",
    "Basic command-line familiarity",
    "Ollama installed (provide install instructions for macOS, Windows, Linux — see Readme.md)",
]
for item in student_prereqs:
    doc.add_paragraph(item, style="List Bullet")

doc.add_heading("Environment Setup Steps (for students)", level=2)
setup_steps = [
    ("Install Ollama", "macOS: brew install ollama | Linux: curl script | Windows: installer download"),
    ("Download model", "ollama pull llama3.2:3b (~2 GB download)"),
    ("Start Ollama", "ollama serve (keep running in a separate terminal)"),
    ("Create venv", "cd FHIRAI && python3 -m venv .venv && source .venv/bin/activate"),
    ("Install deps", "pip install -r requirements.txt"),
    ("Configure", "cp .env.example .env, edit FHIR_PASSWORD and other settings"),
    ("Verify FHIR", "curl http://127.0.0.1:8080/csp/healthshare/demo/fhir/r4/Patient/2"),
]
table = doc.add_table(rows=len(setup_steps) + 1, cols=2)
table.style = "Light Grid Accent 1"
table.rows[0].cells[0].text = "Step"
table.rows[0].cells[1].text = "Command / Action"
for p in table.rows[0].cells[0].paragraphs:
    for r in p.runs:
        r.bold = True
for p in table.rows[0].cells[1].paragraphs:
    for r in p.runs:
        r.bold = True
for i, (step, cmd) in enumerate(setup_steps, start=1):
    table.rows[i].cells[0].text = step
    table.rows[i].cells[1].text = cmd

# ── 4. Phase 1: AI Foundations ──────────────────────────────────────────
doc.add_heading("4. Phase 1: AI Foundations Lecture Guide", level=1)

doc.add_paragraph(
    "This phase covers 9 foundational AI concepts, each grounded in a healthcare analogy. "
    "Full lecture content is in slides/ai_foundations.md. Below are instructor notes for "
    "pacing and emphasis."
)

topics = [
    ("What Is an LLM?",
     "Clinician who has read every textbook but never seen a patient.",
     "Spend time here — this is the conceptual anchor. Students need to understand that "
     "LLMs process text patterns, not structured data. Contrast with FHIR servers."),
    ("Tokens and Tokenization",
     "Pharmacist parsing a prescription into discrete units.",
     "Show the token count example. Students are often surprised that JSON is token-inefficient. "
     "This directly motivates Exercise 2."),
    ("Context Windows",
     "Clinician working memory — can only hold a limited chart in mind at once.",
     "Connect to the concrete numbers from Exercise 2 (~100K tokens for one patient, 8K window). "
     "This is the 'aha moment' that justifies Exercise 3's extraction approach."),
    ("Embeddings",
     "Clinical concept similarity (MI = myocardial infarction = heart attack).",
     "Keep brief — this is a teaser for the RAG Lab. Explain 'meaning as numbers' concept "
     "and move on. Students don't need deep understanding here."),
    ("Temperature and Sampling",
     "Confident specialist (low temp) vs. creative brainstormer (high temp).",
     "Exercise 4 demonstrates this concretely. Preview it: 'You'll see this in action later.'"),
    ("Prompts and System Roles",
     "Clinical protocols that shape every interaction.",
     "Emphasize that system prompts are 'standing orders' for the LLM. Exercises 3 and 4 "
     "both use system prompts — students will write and compare them."),
    ("Hallucination",
     "Clinical confabulation — confident but fabricated.",
     "This is critical for the safety discussion. Give the medication example: LLM invents "
     "a prescription the patient never had. Connects directly to the 'golden rule.'"),
    ("Local vs. Cloud Models",
     "Local lab vs. reference lab.",
     "Emphasize the PHI angle: sending a patient bundle to ChatGPT is a HIPAA violation "
     "without a BAA. This is why we use Ollama."),
    ("Model Selection Basics",
     "Size vs. capability trade-offs.",
     "Quick coverage. Students should know that 3B models are fast but limited, cloud models "
     "are powerful but raise privacy concerns. Our exercises use 3B because it's sufficient "
     "for summarization and runs on any laptop."),
]

for topic, analogy, notes in topics:
    doc.add_heading(topic, level=3)
    p = doc.add_paragraph()
    run = p.add_run("Healthcare analogy: ")
    run.bold = True
    p.add_run(analogy)
    p = doc.add_paragraph()
    run = p.add_run("Instructor notes: ")
    run.bold = True
    p.add_run(notes)

# ── 5. Phase 2: Exercises ───────────────────────────────────────────────
doc.add_heading("5. Phase 2: Exercise Walkthroughs", level=1)

doc.add_paragraph(
    "Each exercise builds on the previous one. Students should complete them in order. "
    "Allow time between exercises for questions and discussion."
)

# Exercise 1
doc.add_heading("5.1 Exercise 1 — LLM Hello World", level=2)
doc.add_paragraph("File: exercise1_llm_chat.py")

doc.add_heading("Purpose", level=3)
doc.add_paragraph(
    "Verify that the student's local LLM setup is working. This is a connectivity test "
    "disguised as a healthcare interaction. The LLM responds to a prompt about A1c testing."
)

doc.add_heading("Running the Exercise", level=3)
doc.add_paragraph("python exercise1_llm_chat.py")

doc.add_heading("What to Expect", level=3)
doc.add_paragraph(
    "The script sends a system prompt ('You are a clinical AI assistant') and a user prompt "
    "('Explain what an A1c test measures, in language suitable for a patient') to the local LLM. "
    "Students should see a multi-sentence response explaining A1c in plain language."
)

doc.add_heading("Common Issues", level=3)
issues_ex1 = [
    "ConnectionError: Ollama is not running. Fix: ollama serve in another terminal.",
    "Model not found: Model not downloaded. Fix: ollama pull llama3.2:3b",
    "Slow response: First call after model load is slow (loading into memory). Subsequent calls are faster.",
]
for issue in issues_ex1:
    doc.add_paragraph(issue, style="List Bullet")

doc.add_heading("Teaching Moment", level=3)
doc.add_paragraph(
    "Point out the system prompt and user prompt in the code. Ask: 'What happens if we change "
    "the system prompt to say You are a pediatric nurse instead?' This previews Exercise 4's "
    "prompt engineering comparison."
)

# Exercise 2
doc.add_heading("5.2 Exercise 2 — FHIR Bundle Exploration", level=2)
doc.add_paragraph("File: exercise2_fhir_explore.py")

doc.add_heading("Purpose", level=3)
doc.add_paragraph(
    "Demonstrate the 'too big for the context window' problem. Students see that a single "
    "patient's $everything bundle is ~100K tokens — far exceeding an 8K context window. "
    "This motivates the extraction approach in Exercise 3."
)

doc.add_heading("Running the Exercise", level=3)
doc.add_paragraph(
    "python exercise2_fhir_explore.py\n"
    "python exercise2_fhir_explore.py --patient 5  (different patient)"
)

doc.add_heading("Key Numbers to Highlight", level=3)
numbers = [
    "Total characters: ~310,000 for a typical Synthea patient",
    "Estimated tokens: ~101,000 (using tiktoken cl100k_base approximation)",
    "Total resources: ~224 (varies by patient)",
    "Context window ratio: ~1,263% of an 8K window — the data is 12x too large",
]
for num in numbers:
    doc.add_paragraph(num, style="List Bullet")

doc.add_heading("Teaching Moment", level=3)
doc.add_paragraph(
    "Ask students: 'If the data doesn't fit, what are our options?' Expected answers: "
    "(1) use a larger model with a bigger context window (expensive, cloud, privacy concerns), "
    "(2) summarize or chunk the data (Exercise 3's approach), "
    "(3) only retrieve relevant parts (RAG Lab's approach). "
    "All three are valid — this lesson teaches option 2, the RAG Lab teaches option 3."
)

# Exercise 3
doc.add_heading("5.3 Exercise 3 — Patient Summarization Pipeline", level=2)
doc.add_paragraph("File: exercise3_summarizer.py")

doc.add_heading("Purpose", level=3)
doc.add_paragraph(
    "The core exercise of the lesson. Students build a complete pipeline: fetch a FHIR bundle, "
    "group resources by type, extract clinically relevant features, send the reduced data to "
    "the LLM, and receive a clinical narrative summary. This achieves ~99% token reduction."
)

doc.add_heading("Running the Exercise", level=3)
doc.add_paragraph(
    "python exercise3_summarizer.py\n"
    "python exercise3_summarizer.py --patient 5  (different patient)"
)

doc.add_heading("The Five Phases", level=3)
phases = [
    ("Phase 1 — Fetch", "Retrieves Patient/$everything bundle from the FHIR server. Same call as Exercise 2."),
    ("Phase 2 — Group", "Organizes the bundle's resources by type (Condition, Observation, MedicationRequest, etc.). This is the key insight: instead of processing raw JSON, we work with structured categories."),
    ("Phase 3 — Extract", "Pulls clinically relevant fields from each resource type: patient demographics, active conditions with onset dates, recent observations with values and units, current medications with dosages, recent encounters with reasons, and allergy intolerances."),
    ("Phase 4 — Summarize", "Sends the extracted features (~1K tokens) to the LLM with a clinical system prompt that includes a 'do not invent information' safety instruction."),
    ("Phase 5 — Output", "Prints the LLM's clinical narrative and reports the token reduction statistics."),
]
for phase, desc in phases:
    p = doc.add_paragraph()
    run = p.add_run(phase + ": ")
    run.bold = True
    p.add_run(desc)

doc.add_heading("Teaching Moments", level=3)
teaching_ex3 = [
    "Token reduction: Point out the ~99% reduction (101K to ~1K tokens). Ask: 'What did we lose? What did we keep? Is the trade-off acceptable for clinical use?'",
    "The 'do not invent' instruction: Show the system prompt. Ask: 'Why is this instruction necessary? What could go wrong without it?' Connect to the hallucination discussion.",
    "Resource-type grouping: This is the same pattern used in production FHIR+AI systems. The student is learning a real technique, not a toy example.",
    "Validation opportunity: Have students compare the LLM summary against the extracted features. Does the summary accurately reflect the data? Are there any hallucinated details?",
]
for item in teaching_ex3:
    doc.add_paragraph(item, style="List Bullet")

doc.add_heading("Discussion Question", level=3)
doc.add_paragraph(
    "'Could a clinician use this summary directly in patient care? What would need to change?' "
    "Expected answer: No — it requires human verification against source data. This is the "
    "'golden rule' introduced in the safety discussion."
)

# Exercise 4
doc.add_heading("5.4 Exercise 4 — Prompt Engineering Comparison", level=2)
doc.add_paragraph("File: exercise4_prompt_compare.py")

doc.add_heading("Purpose", level=3)
doc.add_paragraph(
    "Make prompt engineering concrete and observable. Students see the same patient data "
    "produce visibly different outputs based on prompt structure, system role, examples, "
    "and temperature settings."
)

doc.add_heading("Running the Exercise", level=3)
doc.add_paragraph("python exercise4_prompt_compare.py")
doc.add_paragraph(
    "Note: This script makes 6 LLM calls (3 strategies + 3 temperature variants). "
    "Total runtime may be 1-3 minutes on slower hardware. Progress messages are printed "
    "so students know the script is working."
)

doc.add_heading("The Three Strategies", level=3)

strategies = [
    ("Strategy 1 — Minimal Prompt",
     "Generic 'You are a helpful assistant' with 'Summarize this patient data.' No structure, no role, no safety instructions.",
     "Typically produces an unstructured paragraph. Quality varies. May hallucinate."),
    ("Strategy 2 — Structured System Prompt",
     "Clinical role, explicit section format (Demographics, Conditions, Medications, etc.), 'do not invent' safety instruction.",
     "Produces organized sections matching the requested format. More predictable and safer."),
    ("Strategy 3 — Few-Shot Prompt",
     "Includes an example input/output pair before the real patient data.",
     "Mimics the example's format and style without being told explicitly. Demonstrates learning from examples."),
]
for name, desc, result in strategies:
    p = doc.add_paragraph()
    run = p.add_run(name + ": ")
    run.bold = True
    p.add_run(desc)
    p = doc.add_paragraph()
    run = p.add_run("Expected result: ")
    run.italic = True
    p.add_run(result)

doc.add_heading("Temperature Comparison", level=3)
doc.add_paragraph(
    "Strategy 2 is run three times at different temperatures (0.1, 0.7, 1.0). "
    "Students observe that temperature 0.1 produces nearly identical output each time "
    "(deterministic), while temperature 1.0 introduces more variation. For clinical "
    "summarization, lower temperature is safer."
)

doc.add_heading("Teaching Moments", level=3)
teaching_ex4 = [
    "Have students read the prompt constants at the top of the file. They are designed to be edited — encourage students to modify a prompt and re-run.",
    "Ask: 'Which strategy would you choose for a clinical summarization tool deployed in a hospital?' Expected answer: Strategy 2 (structured, safe, predictable).",
    "Ask: 'When might Strategy 1 (minimal) be acceptable?' Expected answer: Exploration, brainstorming, non-clinical use cases.",
    "Ask: 'What is the risk of high temperature in clinical contexts?' Expected answer: More hallucination risk, less reproducible output.",
]
for item in teaching_ex4:
    doc.add_paragraph(item, style="List Bullet")

# ── 6. Phase 3: Safety ──────────────────────────────────────────────────
doc.add_heading("6. Phase 3: Safety, Governance, and Series Bridge", level=1)

doc.add_paragraph(
    "Full lecture content is in slides/safety_governance.md and slides/series_bridge.md. "
    "This phase frames safety as a series-long theme, not a one-time topic."
)

doc.add_heading("Key Safety Principles to Emphasize", level=2)

safety_points = [
    ("The Golden Rule",
     "Never trust LLM output without validating it against source FHIR data. "
     "This is the single most important safety principle in the series. Students saw this "
     "in Exercise 3 — ask them to recall whether the LLM summary exactly matched the extracted features."),
    ("Hallucination Risks",
     "LLMs can invent medications, fabricate lab values, or conflate patients. "
     "Use the concrete examples from the slides: LLM invents Metformin for a diabetic patient "
     "who was never prescribed it. A clinician acting on this could cause a medication error."),
    ("Human-in-the-Loop",
     "AI assists decisions; clinicians make them. The resident/attending analogy works well: "
     "an attending doesn't sign off on a case summary without reviewing the chart."),
    ("PHI and HIPAA",
     "Copying a patient bundle into ChatGPT is a HIPAA violation. This is why we use Ollama. "
     "Cloud APIs are acceptable only with a BAA in place (e.g., Azure OpenAI Service)."),
    ("Regulatory Awareness",
     "ONC HTI-1 and FDA AI/ML as medical devices — awareness level only. Students should know "
     "these frameworks exist. Do not go deep into compliance requirements."),
]
for title, desc in safety_points:
    p = doc.add_paragraph()
    run = p.add_run(title + ": ")
    run.bold = True
    p.add_run(desc)

doc.add_heading("Series Bridge", level=2)
doc.add_paragraph(
    "The bridge content (slides/series_bridge.md) connects Lesson 17 to the rest of the series. "
    "Cover the four key connection points:"
)
bridge_points = [
    "Summarization pipeline becomes a subagent in Part One",
    "Structured extraction feeds the RAG Lab's vector pipeline",
    "Prompt engineering skills apply to agent prompts in Parts One/Two",
    "Manual FHIR calls (requests) are replaced by MCP in Parts One/Two",
]
for item in bridge_points:
    doc.add_paragraph(item, style="List Bullet")

# ── 7. Timing Guide ─────────────────────────────────────────────────────
doc.add_heading("7. Timing and Pacing Guide", level=1)

doc.add_paragraph(
    "Suggested pacing for a 3-hour session. Adjust based on student experience level "
    "and class engagement."
)

timing = [
    ("Phase 1: AI Foundations Lecture", "45-60 min", "Cover all 9 topics. Spend more time on tokens, context windows, and hallucination — these connect directly to the exercises."),
    ("Setup Check", "10-15 min", "Verify all students have Ollama running and can reach the FHIR server. Troubleshoot connectivity issues."),
    ("Exercise 1: LLM Hello World", "10 min", "Quick smoke test. Move on once students see output."),
    ("Exercise 2: FHIR Bundle Exploration", "15 min", "Run, discuss the numbers, ask 'what are our options?'"),
    ("Exercise 3: Patient Summarization", "30-40 min", "Core exercise. Allow time for discussion of output quality and validation."),
    ("Exercise 4: Prompt Engineering", "20-30 min", "Run, compare outputs, discuss strategy selection."),
    ("Break", "10 min", ""),
    ("Phase 3: Safety Discussion", "20-30 min", "Cover golden rule, hallucination, PHI. Engage students."),
    ("Phase 3: Series Bridge", "10-15 min", "Preview what comes next. Build anticipation."),
    ("Q&A and Wrap-up", "10 min", ""),
]

table = doc.add_table(rows=len(timing) + 1, cols=3)
table.style = "Light Grid Accent 1"
table.rows[0].cells[0].text = "Activity"
table.rows[0].cells[1].text = "Duration"
table.rows[0].cells[2].text = "Notes"
for p in table.rows[0].cells[0].paragraphs:
    for r in p.runs:
        r.bold = True
for p in table.rows[0].cells[1].paragraphs:
    for r in p.runs:
        r.bold = True
for p in table.rows[0].cells[2].paragraphs:
    for r in p.runs:
        r.bold = True
for i, (activity, duration, notes) in enumerate(timing, start=1):
    table.rows[i].cells[0].text = activity
    table.rows[i].cells[1].text = duration
    table.rows[i].cells[2].text = notes

# ── 8. Troubleshooting ──────────────────────────────────────────────────
doc.add_heading("8. Common Issues and Troubleshooting", level=1)

issues = [
    ("Ollama not running",
     "Students get ConnectionError in Exercise 1.",
     "Run 'ollama serve' in a separate terminal. On macOS, Ollama may auto-start as a background service."),
    ("Model not downloaded",
     "Students get 'model not found' error.",
     "Run 'ollama pull llama3.2:3b'. Verify with 'ollama list'."),
    ("FHIR server unreachable",
     "Exercises 2-4 fail to fetch patient data.",
     "Verify IRIS is running. Check FHIR_BASE_URL in .env. Test with curl."),
    ("Wrong Python version",
     "Import errors or syntax issues.",
     "Verify Python 3.11+: python3 --version. Some students may have Python 3.9."),
    ("Missing .env file",
     "Warning about defaults on startup.",
     "Copy .env.example to .env. The defaults work for standard IRIS setup, but password must be set."),
    ("Slow LLM responses",
     "Exercises take several minutes per LLM call.",
     "First call is slow (model loading). Subsequent calls faster. Exercise 4 makes 6 calls — warn students. "
     "Students with 8GB RAM may experience slowness with the 3B model."),
    ("Different output each run",
     "Students get different summaries on the same patient.",
     "This is expected behavior — LLMs are stochastic. Exercise 4 demonstrates this with temperature comparison. "
     "Lower temperature reduces variability."),
]

for title, symptom, fix in issues:
    doc.add_heading(title, level=3)
    p = doc.add_paragraph()
    run = p.add_run("Symptom: ")
    run.bold = True
    p.add_run(symptom)
    p = doc.add_paragraph()
    run = p.add_run("Fix: ")
    run.bold = True
    p.add_run(fix)

# ── 9. Discussion Questions ─────────────────────────────────────────────
doc.add_heading("9. Discussion Questions and Assessment Ideas", level=1)

doc.add_heading("Discussion Questions (during or after exercises)", level=2)
questions = [
    "Exercise 2: 'If the patient bundle is 12x too large for the context window, what information would you prioritize keeping? What would you discard?' (Connects to clinical judgment about what matters.)",
    "Exercise 3: 'The LLM summary says the patient has 3 active conditions. How would you verify this? What if the LLM missed one?' (Introduces validation thinking.)",
    "Exercise 3: 'Could this summarization pipeline run automatically for every patient admission? What safeguards would you need?' (Connects to production deployment, human-in-the-loop.)",
    "Exercise 4: 'A hospital wants to deploy Strategy 2 in their EHR. What additional testing and validation would be required beyond what we did today?' (Connects to regulatory awareness, FDA.)",
    "Safety: 'A developer copies a patient bundle into ChatGPT to test a prompt. What just happened from a HIPAA perspective?' (PHI disclosure, need for BAA.)",
    "Safety: 'If you could only remember one principle from today's safety discussion, what would it be?' (Golden rule: never trust LLM output without validation.)",
    "Series Bridge: 'How would you modify Exercise 3's pipeline to handle 100 patients instead of 1?' (Introduces the need for agents and orchestration — Part One preview.)",
]
for q in questions:
    doc.add_paragraph(q, style="List Bullet")

doc.add_heading("Assessment Ideas", level=2)
assessments = [
    "Have students modify Exercise 3's system prompt and compare outputs — submit the modified prompt and a paragraph explaining what changed and why.",
    "Ask students to identify a hallucination in an LLM output (provide a pre-generated summary with deliberate errors alongside the source FHIR data).",
    "Short reflection: 'Describe one way that AI safety principles from this lesson would apply to a real clinical workflow you've encountered or studied.'",
    "Code modification: Add a new resource type extractor to Exercise 3 (e.g., Procedure or Immunization) and show the updated summary.",
]
for item in assessments:
    doc.add_paragraph(item, style="List Bullet")

# ── 10. Appendix ────────────────────────────────────────────────────────
doc.add_heading("10. Appendix: File Reference", level=1)

files = [
    ("exercise1_llm_chat.py", "Exercise 1: LLM connectivity hello world"),
    ("exercise2_fhir_explore.py", "Exercise 2: FHIR bundle size and token analysis"),
    ("exercise3_summarizer.py", "Exercise 3: Patient summarization CLI pipeline"),
    ("exercise4_prompt_compare.py", "Exercise 4: Prompt engineering comparison"),
    ("config.py", "Configuration loader (reads .env)"),
    ("fhir_client.py", "FHIR server client functions"),
    (".env.example", "Configuration template"),
    ("requirements.txt", "Python dependencies"),
    ("slides/ai_foundations.md", "Phase 1 lecture content (9 AI topics with healthcare analogies)"),
    ("slides/safety_governance.md", "Phase 3 safety and governance content"),
    ("slides/series_bridge.md", "Phase 3 series bridge content (connection to Parts One/Two/RAG Lab)"),
    ("100Set/", "115 Synthea-generated FHIR Bundle JSON files"),
    ("archive/", "Prior development files preserved for reference (not active)"),
]

table = doc.add_table(rows=len(files) + 1, cols=2)
table.style = "Light Grid Accent 1"
table.rows[0].cells[0].text = "File"
table.rows[0].cells[1].text = "Purpose"
for p in table.rows[0].cells[0].paragraphs:
    for r in p.runs:
        r.bold = True
for p in table.rows[0].cells[1].paragraphs:
    for r in p.runs:
        r.bold = True
for i, (fname, purpose) in enumerate(files, start=1):
    table.rows[i].cells[0].text = fname
    table.rows[i].cells[1].text = purpose

# ── Save ────────────────────────────────────────────────────────────────
output_path = os.path.join(os.path.dirname(__file__), "Lesson17_Instructor_Guide.docx")
doc.save(output_path)
print(f"Saved: {output_path}")
